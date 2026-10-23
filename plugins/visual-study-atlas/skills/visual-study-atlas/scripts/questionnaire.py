#!/usr/bin/env python3
"""Questionnaire for visual-study-atlas (max 25 questions).
  questionnaire.py make context.json questions_<name>.xlsx
  questionnaire.py read questions_<name>.xlsx answers.json
context.json: {"topic": "...", "material": "...", "pilot": "what the pilot covered",
  "assumed": {question_id: "assumption text"}, "drop": [ids], "extra": [{"id","question","choices":[..],"assumed":".."}] (max 3)}"""
import sys, json
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, Alignment, PatternFill

CORE = [
 ("material", "What is the material?", ["book", "research paper", "article / blog", "video or lecture transcript", "code repo / notebook", "course (many modules)", "mixed notes"]),
 ("purpose", "What is the main purpose?", ["pass an exam", "do a job role better", "teach others", "general understanding", "research / reproduce", "interview preparation"]),
 ("target", "Which exam, role or audience exactly? (free text)", None),
 ("prior", "Your starting knowledge of this topic?", ["none", "some", "strong"]),
 ("depth", "How deep?", ["overview", "working understanding", "deep with derivations"]),
 ("focus", "What should the pages center on?", ["core equations first", "pipeline and architecture first", "balanced equations and pipelines", "concepts and definitions", "argument and evidence", "you decide"]),
 ("maths", "How much mathematics?", ["none, plain words", "formulas explained with symbols", "step-by-step derivations", "full derivations and proofs"]),
 ("code", "How much code?", ["none", "key snippets", "function-by-function with how they link", "full walkthrough"]),
 ("scope", "Scope?", ["everything", "chosen chapters / sections (list in comments)", "one topic only"]),
 ("include", "Must include (free text)", None),
 ("exclude", "Leave out (free text)", None),
 ("visuals", "Which pictures help you most?", ["flow and architecture diagrams", "equation maps and plots", "tables and matrices", "timelines and cycles", "mix, you decide"]),
 ("examples", "Worked examples?", ["hypothetical, small numbers", "real data or case studies from the source", "both", "none"]),
 ("lenses", "Closing panels per section?", ["none", "exam lens", "professional lens", "both", "custom (name them in comments)"]),
 ("selfcheck", "Self-check questions?", ["none", "a few per chapter", "full set with worked answers"]),
 ("quizstyle", "Quiz style for the whole project? (scored, per chapter)", ["certification-exam style: scenario, pick the BEST of 4, explain every wrong option", "plain recall: term or concept to meaning, 4 options", "true / false with a reason", "mixed", "none"]),
 ("features", "Study extras to include?", ["all: page-cited ask box, flashcard export, mind map, step-by-step flow animations", "ask box and flashcard export only", "none of these"]),
 ("askbox", "Click-anywhere ask box?", ["yes", "no"]),
 ("askstyle", "Ask-box answer style?", ["very compact (under 60 words)", "explainer (110 to 170 words, mechanism + example)", "step-by-step tutor"]),
 ("delivery", "Where should the final version live? (asked every time)", ["local folder (offline, runs from Finder)", "private published artifact (link)", "both"]),
 ("videolinks", "Add reference videos (verified YouTube links) for technical or dense topics?", ["yes, only where the topic is dense", "yes, for every section", "no"]),
 ("sourcelinks", "Link each section back to the source (page / section / line) so you can read side by side?", ["yes, to pages of my copy", "yes, to section titles only", "no"]),
 ("engine", "Who answers the ask box?", ["local Ollama models (free, offline)", "Claude (best quality, uses your usage)", "no ask box"]),
 ("look", "Look, feel and tone?", ["clean light and dark, airy, plain friendly English", "dense notes style, plain formal English", "airy, plain English with simple glosses in another language (say which in comments)", "you decide"]),
 ("pilot", "What would you change in the pilot?", None),
]

def make(ctx, out):
    drop = set(ctx.get("drop", [])); assumed = ctx.get("assumed", {})
    qs = [q for q in CORE if q[0] not in drop]
    for e in ctx.get("extra", [])[:3]: qs.append((e["id"], e["question"], e.get("choices")))
    qs = qs[:25]
    wb = Workbook(); ws = wb.active; ws.title = "Questions"
    ws.append(["#", "Question", "What the pilot assumed", "Your answer (pick or type)", "Comments"])
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0B7A75"); c.alignment = Alignment(wrap_text=True)
    ids = []
    for i, (qid, q, ch) in enumerate(qs, 1):
        ws.append([i, q, assumed.get(qid, ""), "", ""]); ids.append(qid)
        r = i + 1
        if ch:
            dv = DataValidation(type="list", formula1='"' + ",".join(x.replace(",", ";") for x in ch) + '"', allow_blank=True); dv.error = "You may also type your own answer"; dv.errorStyle = "warning"
            ws.add_data_validation(dv); dv.add(f"D{r}")
    for col, w in zip("ABCDE", (4, 52, 36, 40, 50)): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ov = wb.create_sheet("Overall"); ov["A1"] = "Anything unique you need (extra requirement, format, constraint)"; ov["A1"].font = Font(bold=True)
    ov["A2"] = ""; ov.column_dimensions["A"].width = 110; ov["A2"].alignment = Alignment(wrap_text=True, vertical="top"); ov.row_dimensions[2].height = 160
    meta = wb.create_sheet("meta"); meta["A1"] = json.dumps({"ids": ids, "topic": ctx.get("topic", ""), "pilot": ctx.get("pilot", "")}); meta.sheet_state = "hidden"
    wb.save(out); print(f"wrote {out} with {len(qs)} questions")

def read(path, out):
    wb = load_workbook(path); ws = wb["Questions"]; meta = json.loads(wb["meta"]["A1"].value); ids = meta["ids"]; ans = {}; comments = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if i >= len(ids): break
        _, q, assumed, a, c = (list(row) + [None] * 5)[:5]
        ans[ids[i]] = {"question": q, "assumed": assumed or "", "answer": (a or assumed or "").strip() if isinstance(a or assumed or "", str) else a, "answered": bool(a)}
        if c: comments.append({"id": ids[i], "comment": str(c)})
    overall = wb["Overall"]["A2"].value or ""
    json.dump({"topic": meta.get("topic"), "answers": ans, "comments": comments, "overall": overall}, open(out, "w"), indent=1)
    print(f"read {len(ans)} answers, {len(comments)} comments")

if __name__ == "__main__":
    if sys.argv[1] == "make": make(json.load(open(sys.argv[2])), sys.argv[3])
    elif sys.argv[1] == "read": read(sys.argv[2], sys.argv[3])
    else: print(__doc__)
